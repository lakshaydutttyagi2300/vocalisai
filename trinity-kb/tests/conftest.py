import pytest


@pytest.fixture
def conn(tmp_path, monkeypatch):
    monkeypatch.setenv("TKB_DATA_DIR", str(tmp_path / "data"))
    from tkb import db
    c = db.connect()
    yield c
    c.close()


@pytest.fixture
def files(tmp_path):
    """Synthetic test documents (not real Trinity content)."""
    import docx
    import pymupdf as fitz
    import openpyxl
    d = tmp_path / "in"; d.mkdir()
    pdf = fitz.open()
    p1 = pdf.new_page(); p1.insert_text((72, 72), "Exam Overview", fontsize=18)
    p1.insert_text((72, 110), "Candidates prepare three pieces for the exam.", fontsize=11)
    p2 = pdf.new_page(); p2.insert_text((72, 72), "Digital Submissions", fontsize=18)
    p2.insert_text((72, 110), "After a digital exam recording is submitted, results are issued within 15 working days.", fontsize=11)
    pdf.save(d / "Syllabus.pdf")
    doc = docx.Document()
    doc.add_heading("Face-to-face bookings", 1)
    doc.add_paragraph("Tickets about a face-to-face booking go to the Bookings team.")
    doc.add_heading("Replacement certificates", 1)
    doc.add_paragraph("To request a replacement certificate the candidate completes the replacement form.")
    t = doc.add_table(rows=2, cols=3)
    for i, v in enumerate(["Issue", "Group", "Priority"]): t.rows[0].cells[i].text = v
    for i, v in enumerate(["LAR code used by teacher", "Centre Support", "Normal"]): t.rows[1].cells[i].text = v
    doc.save(d / "FCR_Triage_Directory.docx")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "General"
    ws.append(["Question", "Answer"])
    ws.append(["Which profile type should a parent select?", "A parent should select the Parent/Guardian profile type."])
    ws2 = wb.create_sheet("Digital")
    ws2.append(["Question", "Answer"])
    ws2.append(["How is a digital booking confirmed?", "The candidate confirms by clicking the link in the confirmation email. Booking references look like DG1234567."])
    ws3 = wb.create_sheet("Old"); ws3.sheet_state = "hidden"
    ws3.append(["Note"]); ws3.append(["Results for digital recordings are issued within 20 working days."])
    wb.save(d / "Q&A.xlsx")
    return d
