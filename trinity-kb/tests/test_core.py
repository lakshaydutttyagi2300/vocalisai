import re
import json

from tkb import backup, db, ingest, ticket, web
from tkb.search import citation, search


def imp(conn, files):
    return {f.name: ingest.import_file(conn, f) for f in sorted(files.iterdir())}


# 1-3 importing
def test_import_pdf_docx_xlsx(conn, files):
    r = imp(conn, files)
    assert all(x["status"] == "added" for x in r.values()), r
    pdf = conn.execute("SELECT * FROM passages p JOIN sources s ON s.id=p.source_id WHERE s.name='Syllabus.pdf'").fetchall()
    assert {p["page"] for p in pdf} == {1, 2}
    sheets = {p["sheet"] for p in conn.execute("SELECT sheet FROM passages WHERE sheet IS NOT NULL")}
    assert sheets == {"General", "Digital", "Old"}  # including the hidden sheet
    warn = json.loads(conn.execute("SELECT warnings FROM sources WHERE name='Q&A.xlsx'").fetchone()[0])
    assert any("hidden" in w for w in warn)
    docx_rows = [p["text"] for p in conn.execute("SELECT text FROM passages p JOIN sources s ON s.id=p.source_id WHERE s.name LIKE '%.docx'")]
    assert any("Group: Centre Support" in t for t in docx_rows)  # table kept with its headers
    assert conn.execute("SELECT category FROM sources WHERE name='Q&A.xlsx'").fetchone()[0] == "D"


# 4 duplicates
def test_duplicate_document(conn, files, tmp_path):
    imp(conn, files)
    assert ingest.import_file(conn, files / "Syllabus.pdf")["status"] == "unchanged"
    copy = tmp_path / "Copy of syllabus.pdf"; copy.write_bytes((files / "Syllabus.pdf").read_bytes())
    assert ingest.import_file(conn, copy)["status"] == "duplicate"
    assert conn.execute("SELECT COUNT(*) FROM sources").fetchone()[0] == 3


# 5 updates
def test_changed_document_updates_index(conn, files):
    import docx
    imp(conn, files)
    d = docx.Document(); d.add_heading("Face-to-face bookings", 1); d.add_paragraph("Face-to-face booking tickets now go to the Regional team.")
    d.save(files / "FCR_Triage_Directory.docx")
    r = ingest.import_file(conn, files / "FCR_Triage_Directory.docx")
    assert r["status"] == "updated"
    assert search(conn, "Regional team")
    assert not search(conn, '"Bookings team"')
    s = conn.execute("SELECT version FROM sources WHERE name='FCR_Triage_Directory.docx'").fetchone()
    assert s[0] == 2 and conn.execute("SELECT COUNT(*) FROM source_versions WHERE source_id=(SELECT id FROM sources WHERE name='FCR_Triage_Directory.docx')").fetchone()[0] == 2


def test_corrupt_file_reported(conn, tmp_path):
    bad = tmp_path / "broken.pdf"; bad.write_bytes(b"not a pdf at all")
    r = ingest.import_file(conn, bad)
    assert r["status"] == "failed"
    assert conn.execute("SELECT status FROM sources WHERE name='broken.pdf'").fetchone()[0] == "failed"


# 6 phrase search, 8/9 correct passage + citation
def test_phrase_search_and_citation(conn, files):
    imp(conn, files)
    res = search(conn, '"digital exam recording is submitted"')
    assert len(res) == 1
    r = res[0]
    assert r["name"] == "Syllabus.pdf" and r["page"] == 2
    c = citation(r)
    assert "page 2" in c["text"] and "Digital Submissions" in c["text"]
    assert search(conn, '"recording is digital exam"') == [] or all("recording is digital" not in x["text"] for x in search(conn, '"recording is digital exam"'))


def test_docx_and_xlsx_citations(conn, files):
    imp(conn, files)
    r = search(conn, "replacement certificate")[0]
    assert "section 'Replacement certificates'" in citation(r)["text"]
    r = search(conn, "profile type parent")[0]
    c = citation(r)["text"]
    assert "worksheet 'General'" in c and "cells A2:B2" in c


# 7 booking-reference format
def test_reference_format_search(conn, files):
    imp(conn, files)
    res = search(conn, "format:DG7654321")  # different digits, same shape
    assert res and "DG1234567" in res[0]["text"]
    assert search(conn, "DG1234567")[0]["sheet"] == "Digital"


def test_email_and_filters(conn, files, tmp_path):
    f = tmp_path / "contacts.txt"; f.write_text("Escalations go to support.team@example.org for review.")
    ingest.import_file(conn, f, "B")
    imp(conn, files)
    assert search(conn, "support.team@example.org")[0]["name"] == "contacts.txt"
    assert all(r["category"] == "D" for r in search(conn, "digital", categories=["D"]))
    assert all(r["heading"] and "Digital" in r["heading"] for r in search(conn, "digital", heading="Digital"))


# 10 missing evidence, 12 no invention, 13 no AI
def test_missing_evidence_and_no_model(conn, files):
    imp(conn, files)
    res = ticket.run_assistant(conn, "Hello, can I get a refund for my parking fine at the venue car park?", "")
    assert res["assessment"]["verdict"] == "insufficient"
    assert ticket.INSUFFICIENT in res["draft"] and ticket.INSUFFICIENT in res["warnings"][0]
    assert res["mode"] == "retrieval-only"
    assert res["draft"].startswith(ticket.OPENING)


def test_answer_uses_only_evidence(conn, files):
    imp(conn, files)
    res = ticket.run_assistant(conn, "Hi, my daughter booked a digital exam. How does she confirm the digital booking?", "How is a digital booking confirmed?")
    assert res["assessment"]["verdict"] in ("sufficient", "partial")
    body = [l for l in res["draft"].splitlines() if re.search(r"\[E\d+\]$", l)]
    assert body, res["draft"]
    pool = " ".join(e["text"] for e in res["evidence"])
    for line in body:  # every evidence sentence is verbatim from a source
        sentence = line.lstrip("- ").rsplit(" [E", 1)[0]
        assert sentence in pool
    assert "parent / guardian" in res["analysis"]["role"]
    assert res["link"] is None  # booking question: no syllabus link


def test_generated_text_guard():
    ev = [{"text": "Results are issued within 15 working days."}]
    out, warns = ticket.check_generated("Your refund of £50 will arrive in 3 days. See https://fake.example/x", ev, "", None)
    assert out.startswith(ticket.OPENING) and "fake.example" not in out
    assert any("£50" in w for w in warns) and any("'3'" in w for w in warns)


# 11 conflicts
def test_conflict_detection(conn, files):
    imp(conn, files)
    res = ticket.run_assistant(conn, "When are results for a digital recording issued after it is submitted?", "")
    assert res["conflicts"], [e["text"] for e in res["evidence"]]
    cf = res["conflicts"][0]
    assert {"15", "20"} == set(cf["va"]) | set(cf["vb"])
    assert any("Conflicting" in w for w in res["warnings"])


def test_triage_uses_internal_sources(conn, files):
    imp(conn, files)
    t = ticket.triage(conn, "A teacher used an LAR code to book, which group handles this?")
    top = t["recommendations"][0]
    assert ("Group", "Centre Support") in top["fields"]
    assert top["row"]["category"] in "BCD"
    empty = ticket.triage(conn, "Question about the colour of the building")
    assert empty["certainty"] == "uncertain"


# 14 offline after indexing
def test_works_offline(conn, files, monkeypatch):
    imp(conn, files)
    import socket

    def no_net(*a, **k):
        raise OSError("network disabled")
    monkeypatch.setattr(socket, "create_connection", no_net)
    monkeypatch.setattr(socket.socket, "connect", no_net)
    assert search(conn, "replacement certificate")
    res = ticket.run_assistant(conn, "How do I get a replacement certificate?", "")
    assert res["evidence"]


# 15 backup and restore
def test_backup_restore(conn, files):
    imp(conn, files)
    z = backup.create_backup(conn)
    n = conn.execute("SELECT COUNT(*) FROM passages").fetchone()[0]
    ingest.remove_source(conn, conn.execute("SELECT id FROM sources WHERE name='Syllabus.pdf'").fetchone()[0])
    conn.close()
    backup.restore_backup(z)
    c2 = db.connect()
    assert c2.execute("SELECT COUNT(*) FROM passages").fetchone()[0] == n
    assert search(c2, '"three pieces"')
    from pathlib import Path
    assert all(Path(r[0]).exists() for r in c2.execute("SELECT stored_path FROM sources"))


# 16 website failures, robots, duplicates, resume
class FakeSite:
    def __init__(self, pages):
        self.pages, self.calls = pages, []

    def get(self, url, headers=None):
        self.calls.append(url)
        if url not in self.pages:
            return 404, {}, b"", url
        st, ct, body = self.pages[url]
        if isinstance(st, Exception):
            raise st
        return st, {"content-type": ct}, body.encode() if isinstance(body, str) else body, url


B = "https://www.trinitycollege.com"


def site():
    page = lambda t, body, links="": f"<html><head><title>{t}</title></head><body><main><h1>{t}</h1><p>{body}</p>{links}</main></body></html>"
    sitemap = f"""<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
      <url><loc>{B}/music</loc></url><url><loc>{B}/music-copy</loc></url><url><loc>{B}/private/staff</loc></url>
      <url><loc>{B}/broken</loc></url><url><loc>{B}/locked</loc></url><url><loc>{B}/gone</loc></url></urlset>"""
    return {
        f"{B}/robots.txt": (200, "text/plain", f"User-agent: *\nDisallow: /private/\nSitemap: {B}/sitemap.xml"),
        f"{B}/sitemap.xml": (200, "application/xml", sitemap),
        f"{B}/": (200, "text/html", page("Home", "Welcome to the music exams information.", f'<a href="/music">m</a><a href="https://evil.example/x">x</a>')),
        f"{B}/music": (200, "text/html", page("Music grades", "Grade 5 theory requirements are described in the syllabus.")),
        f"{B}/music-copy": (200, "text/html", page("Music grades", "Grade 5 theory requirements are described in the syllabus.")),
        f"{B}/broken": (500, "text/html", "err"),
        f"{B}/locked": (403, "text/html", "no"),
    }


def test_web_requires_terms_confirmation(conn):
    import pytest
    with pytest.raises(web.CrawlBlocked):
        web.run_update(conn, FakeSite(site()), sleep=lambda s: None)


def test_web_ingestion_and_coverage(conn):
    db.set_setting(conn, "web_terms_reviewed", True)
    f = FakeSite(site())
    sleeps = []
    r = web.run_update(conn, f, sleep=sleeps.append)
    cov = web.coverage_report(conn)
    assert cov["processed"] == 2 and cov["duplicate"] == 1
    assert cov["failed"] == 2 and cov["manual"] == 1 and cov["skipped"] == 1
    assert f"{B}/private/staff" not in f.calls  # robots.txt respected
    assert not any("evil.example" in u for u in f.calls)
    assert sleeps and min(sleeps) >= 3.0  # rate limit
    states = {p["url"]: p for p in cov["problems"]}
    assert "robots" in states[f"{B}/private/staff"]["note"]
    assert conn.execute("SELECT COUNT(*) FROM job_errors").fetchone()[0] >= 2
    hit = search(conn, "grade 5 theory requirements")[0]
    c = citation(hit)
    assert c["url"] == f"{B}/music" and c["category"] == "A" and c["verified"] != "not recorded"
    assert "Incomplete" not in cov["verdict"] and "does not prove" in cov["verdict"]


def test_web_resume_and_change_detection(conn):
    db.set_setting(conn, "web_terms_reviewed", True)
    pages = site()
    r = web.run_update(conn, FakeSite(pages), sleep=lambda s: None, max_pages=2)
    assert r["pending"] > 0 and web.coverage_report(conn)["verdict"].startswith("Incomplete")
    web.run_update(conn, FakeSite(pages), sleep=lambda s: None)
    assert web.coverage_report(conn)["pending"] == 0
    pages[f"{B}/music"] = (200, "text/html", "<html><title>Music grades</title><main><h1>Music grades</h1><p>Grade 5 theory requirements changed for 2027.</p></main></html>")
    import tkb.db as d
    conn.execute("UPDATE web_pages SET verified_at='2000-01-01' WHERE url=?", (f"{B}/music",)); conn.commit()
    r = web.run_update(conn, FakeSite(pages), sleep=lambda s: None, refresh=True)
    assert r["updated"] >= 1 and search(conn, "changed 2027")
    assert conn.execute("SELECT version FROM sources WHERE name=?", (f"{B}/music",)).fetchone()[0] == 2


def test_web_network_failure_reported(conn):
    db.set_setting(conn, "web_terms_reviewed", True)
    r = web.run_update(conn, FakeSite({f"{B}/robots.txt": (ConnectionError("down"), "", "")}), sleep=lambda s: None)
    assert r["status"] == "failed" and "Nothing was crawled" in r["message"]


def test_official_link_only_from_index(conn):
    db.set_setting(conn, "web_terms_reviewed", True)
    web.run_update(conn, FakeSite(site()), sleep=lambda s: None)
    res = ticket.run_assistant(conn, "What are the grade 5 theory requirements in the syllabus?", "")
    assert res["link"] == f"{B}/music"
    assert res["draft"].count("http") == 1


# 17 duplicate passages
def test_no_duplicate_passages(conn, tmp_path):
    f = tmp_path / "dupes.md"
    f.write_text("# A\nSame paragraph text here.\n# B\nSame paragraph text here.\n# C\nDifferent text.")
    r = ingest.import_file(conn, f)
    assert r["passages"] == 2 and r["duplicates"] == 1


# 18 tickets kept out of the knowledge base
def test_ticket_not_stored(conn, files):
    imp(conn, files)
    before = conn.execute("SELECT COUNT(*) FROM passages").fetchone()[0]
    secret = "My name is Jane Example, email jane@example.com, ref DG5550001"
    ticket.run_assistant(conn, secret + " How do I confirm my digital booking?", "")
    assert conn.execute("SELECT COUNT(*) FROM passages").fetchone()[0] == before
    dump = "\n".join(conn.iterdump())
    assert "jane@example.com" not in dump and "Jane Example" not in dump
    red = ticket.redact("Dear Jane,\nmy email jane@example.com ref DG5550001\nKind regards,\nJane")
    assert "jane@example.com" not in red and "DG5550001" not in red and "Jane" not in red
    ingest.add_note(conn, "Example", red, "D")
    assert conn.execute("SELECT COUNT(*) FROM passages").fetchone()[0] == before + 1


def test_trinity_style_resources_and_noindex(conn):
    """Mirrors the live site: pages tagged noindex link to syllabus PDFs at /resource/?id=NNNN."""
    import pymupdf
    pdf = pymupdf.open(); pg = pdf.new_page()
    pg.insert_text((72, 72), "PIANO SYLLABUS", fontsize=18); pg.insert_text((72, 110), "Candidates perform three pieces and technical work.")
    pdf_bytes = pdf.tobytes()
    noindex = '<html><head><meta name="robots" content="noindex"><title>Piano</title></head><body><main><h1>Piano</h1><p>Piano grade exams overview.</p>'
    pages = {
        f"{B}/robots.txt": (200, "text/plain", f"User-agent: *\nAllow: /\nSitemap: {B}/sitemap.xml"),
        f"{B}/sitemap.xml": (200, "application/xml", f'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>{B}/piano/</loc></url></urlset>'),
        f"{B}/": (200, "text/html", noindex + '</main></body></html>'),
        f"{B}/piano": (200, "text/html", noindex + f'<a href="/resource/?id=10125">Syllabus</a><a href="/search?q=x">s</a></main></body></html>'),
        f"{B}/resource?id=10125": (200, "application/pdf", pdf_bytes),
    }
    db.set_setting(conn, "web_terms_reviewed", True)
    f = FakeSite(pages)
    web.run_update(conn, f, sleep=lambda s: None)
    cov = web.coverage_report(conn)
    assert cov["processed"] == 1 and cov["skipped"] == 2  # PDF kept; both noindex pages skipped
    r = search(conn, "three pieces technical work")[0]
    assert r["location"] == f"{B}/resource?id=10125" and r["category"] == "A" and r["title"] == "PIANO SYLLABUS"
    from pathlib import Path
    assert Path(r["stored_path"]).read_bytes()[:5] == b"%PDF-"  # the official file is saved for offline use
    assert not any("search" in u for u in f.calls)
    db.set_setting(conn, "web_respect_noindex", False)
    conn.execute("UPDATE web_pages SET state='pending' WHERE state='skipped'"); conn.commit()
    web.run_update(conn, FakeSite(pages), sleep=lambda s: None)
    cov = web.coverage_report(conn)
    assert cov["processed"] == 2 and cov["duplicate"] == 1 and search(conn, "piano grade exams overview")


def test_backup_to_drive_folder(conn, files, tmp_path):
    drive = tmp_path / "My Drive"; drive.mkdir()
    imp(conn, files)
    assert backup.auto_backup(conn) is None  # no folder chosen: nothing automatic
    db.set_setting(conn, "backup_folder", str(drive)); db.set_setting(conn, "backup_keep", 2)
    first = backup.auto_backup(conn)
    assert first and first.parent == drive / "TrinityKB-backups"
    assert backup.auto_backup(conn) is None  # at most once a day
    for _ in range(3):
        backup.create_backup(conn)
        import time; time.sleep(1.1)
    db.set_setting(conn, "last_backup_at", "2000-01-01")
    backup.auto_backup(conn)
    assert len(list((drive / "TrinityKB-backups").glob("*.zip"))) == 2  # old copies pruned
